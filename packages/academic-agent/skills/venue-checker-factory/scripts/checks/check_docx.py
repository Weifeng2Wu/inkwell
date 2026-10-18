"""Word (.docx) checks. Layout-dependent rules (page count) need a PDF export: see run_checks --docx-to-pdf."""
import re
from collections import Counter

from docx import Document

from .common import finding, margins_pt, sources, Finding

PAGE_SIZES = {"A4": (595.28, 841.89), "LETTER": (612.0, 792.0)}


def _norm(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def run(path, rules, names=None):
    D = rules.get("docx", {}) or {}
    S = sources(D)
    F = []
    doc = Document(path)
    sec = doc.sections[0]
    ctx = {"text": "\n".join(p.text for p in doc.paragraphs)}

    if D.get("page_size"):
        want = PAGE_SIZES[D["page_size"].upper()]
        w, h = sec.page_width.pt, sec.page_height.pt
        ok = abs(w - want[0]) <= 3 and abs(h - want[1]) <= 3
        F.append(finding("docx.page_size", ok, f"Page size is {D['page_size']}", f"Document is {w:.0f}x{h:.0f} pt",
                         fix="Layout > Size > choose the required paper size.", source=S.get("page_size", "")))

    M = margins_pt(D.get("margins"))
    if M:
        tol = float((D.get("margins") or {}).get("tolerance_pt", 3))
        have = {"top": sec.top_margin.pt, "bottom": sec.bottom_margin.pt,
                "left": sec.left_margin.pt, "right": sec.right_margin.pt}
        bad = [f"{k}: {have[k]:.0f}pt vs required {v:.0f}pt" for k, v in M.items() if abs(have[k] - v) > tol]
        F.append(finding("docx.margins", not bad, "Margins match the requirement", "; ".join(bad),
                         fix="Layout > Margins > Custom Margins, enter the values from the kit.",
                         source=S.get("margins", "")))

    # fonts: count characters per (font, size), falling back to the style / Normal defaults
    normal = doc.styles["Normal"].font
    cnt_font, cnt_size = Counter(), Counter()
    spacing = Counter()
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        pf = p.style.font if p.style is not None else None
        spacing[p.paragraph_format.line_spacing or (p.style.paragraph_format.line_spacing if p.style is not None else None)] += len(p.text)
        for r in p.runs:
            n = len(r.text.strip())
            if not n:
                continue
            name = r.font.name or (pf.name if pf is not None else None) or normal.name
            size = r.font.size or (pf.size if pf is not None else None) or normal.size
            cnt_font[name] += n
            cnt_size[round(size.pt, 1) if size else None] += n
    FO = D.get("fonts") or {}
    allowed = [_norm(a) for a in FO.get("allowed_families", [])]
    if allowed:
        bad = {f: c for f, c in cnt_font.items() if not any(a in _norm(f) for a in allowed)}
        F.append(finding("docx.fonts_allowed", not bad, "Only the required font family is used",
                         "" if not bad else "Other fonts: " + ", ".join(f"{f or 'default'} ({c} chars)" for f, c in bad.items()),
                         fix="Select all (Ctrl+A) and set the font from the kit; check headings, captions and tables.",
                         source=S.get("fonts", "")))
    if FO.get("body_size_pt"):
        dom = cnt_size.most_common(1)[0][0] if cnt_size else None
        F.append(finding("docx.font_size_body", dom is not None and abs(dom - FO["body_size_pt"]) <= 0.35,
                         f"Body text is {FO['body_size_pt']} pt", f"Most common size: {dom} pt",
                         fix="Set the Normal style to the required size instead of formatting runs by hand.",
                         source=S.get("fonts", "")))
    if D.get("line_spacing") is not None:
        dom = spacing.most_common(1)[0][0] if spacing else None
        F.append(finding("docx.line_spacing", dom is not None and abs(float(dom) - D["line_spacing"]) <= 0.05,
                         f"Line spacing is {D['line_spacing']}", f"Dominant line spacing: {dom}",
                         fix="Paragraph > Line spacing: set the value from the kit.", source=S.get("line_spacing", "")))
    if D.get("page_limit"):
        F.append(Finding("docx.page_limit", "manual", f"Page limit ({D['page_limit']})",
                         "Word does not store page breaks reliably. Re-run with --docx-to-pdf (LibreOffice) or export to PDF "
                         "from Word and pass --pdf. Remember the PDF produced by LibreOffice may differ slightly from Word."))
    return F, ctx
