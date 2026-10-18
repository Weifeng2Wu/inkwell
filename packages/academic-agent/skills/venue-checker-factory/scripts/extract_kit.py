#!/usr/bin/env python3
"""Turn an author kit (zip / folder / loose files) into readable text for rule extraction.

  python scripts/extract_kit.py <file-or-dir> [<more>...] --out kit_text/

Writes into --out:
  <name>.txt          text of each pdf / docx / tex / html / md / txt document
  style_digest.md     layout-relevant lines of every .sty/.cls/.bst/.tex (margins, fonts, options, packages)
  inventory.json      every file with size + sha256 (use it for tex.official_files) and a guess at its role
Prints a short inventory so the caller can see what was found.
"""
import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

TEXT_EXT = {".tex", ".sty", ".cls", ".bst", ".txt", ".md", ".bib", ".def", ".cfg", ".rtf"}
STYLE_EXT = {".sty", ".cls", ".bst"}
DIGEST_RE = re.compile(
    r"textwidth|textheight|oddsidemargin|evensidemargin|topmargin|headheight|headsep|footskip|columnsep|"
    r"paperwidth|paperheight|geometry|\\RequirePackage|\\usepackage|DeclareOption|ProcessOptions|ExecuteOptions|"
    r"if(review|final|preprint|anon|camera|submission)|\\if@|fontsize|linespread|baselinestretch|"
    r"rmdefault|\\usepackage\{(times|mathptmx|newtxtext|helvet|txfonts)\}|bibliographystyle|natbib|"
    r"@maketitle|@makefntext|\\thanks|anonymous|Anonymous|\\setlength|\\addtolength|\\@setfontsize|\\renewcommand\{\\normalsize\}",
    re.I)
MAX_DIGEST_LINES_PER_FILE = 220


def sha(path):
    return hashlib.sha256(Path(path).read_bytes().replace(b"\r\n", b"\n")).hexdigest()


def pdf_text(path):
    exe = shutil.which("pdftotext")
    if exe:
        r = subprocess.run([exe, "-layout", str(path), "-"], capture_output=True, text=True, timeout=120)
        if r.stdout.strip():
            return r.stdout
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        return "\n\f\n".join((p.extract_text() or "") for p in pdf.pages)


def docx_text(path):
    from docx import Document
    d = Document(path)
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    sec = d.sections[0]
    out.append(f"\n[page setup] width={sec.page_width.pt:.0f}pt height={sec.page_height.pt:.0f}pt "
               f"margins(top,bottom,left,right)={sec.top_margin.pt:.0f},{sec.bottom_margin.pt:.0f},"
               f"{sec.left_margin.pt:.0f},{sec.right_margin.pt:.0f}pt")
    return "\n".join(out)


def html_text(path):
    t = Path(path).read_text(errors="ignore")
    t = re.sub(r"(?is)<(script|style).*?</\1>", " ", t)
    t = re.sub(r"(?s)<[^>]+>", " ", t)
    return re.sub(r"[ \t]+", " ", re.sub(r"\n\s*\n+", "\n\n", t))


def guess_role(p):
    n = p.name.lower()
    if p.suffix.lower() in STYLE_EXT:
        return "style file (ground truth for layout; hash it)"
    if p.suffix.lower() == ".tex":
        t = p.read_text(errors="ignore")
        return "LaTeX template / sample" if "\\documentclass" in t else "LaTeX fragment"
    if p.suffix.lower() in (".pdf", ".docx", ".doc", ".md", ".txt", ".html", ".htm"):
        if any(k in n for k in ("checklist", "responsible", "reproducib")):
            return "checklist"
        if any(k in n for k in ("author", "guideline", "instruction", "call", "cfp", "format", "template", "kit")):
            return "formatting instructions"
        return "document"
    return "other"


def collect(inputs, work):
    files = []
    for inp in inputs:
        p = Path(inp)
        if p.is_dir():
            files += [f for f in p.rglob("*") if f.is_file()]
        elif p.suffix.lower() == ".zip":
            dest = work / "_unzipped" / p.stem
            dest.mkdir(parents=True, exist_ok=True)
            zipfile.ZipFile(p).extractall(dest)
            files += [f for f in dest.rglob("*") if f.is_file()]
        elif p.is_file():
            files.append(p)
    return [f for f in files if "__MACOSX" not in f.parts and not f.name.startswith(".")]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("inputs", nargs="+")
    ap.add_argument("--out", default="kit_text")
    a = ap.parse_args()
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    files = collect(a.inputs, out)

    inventory, digest = [], ["# Style digest (layout-relevant lines, with file:line)\n"]
    for f in sorted(files):
        ext = f.suffix.lower()
        rec = {"path": str(f), "name": f.name, "bytes": f.stat().st_size, "sha256": sha(f), "role": guess_role(f)}
        txt = None
        try:
            if ext == ".pdf":
                txt = pdf_text(f)
            elif ext == ".docx":
                txt = docx_text(f)
            elif ext in (".html", ".htm"):
                txt = html_text(f)
            elif ext in TEXT_EXT:
                txt = f.read_text(errors="ignore")
        except Exception as e:  # keep going on a bad file
            rec["error"] = str(e)
        if txt is not None:
            tgt = out / (f.name + ".txt")
            tgt.write_text(txt, encoding="utf-8")
            rec["text_file"] = str(tgt)
            rec["chars"] = len(txt)
            if ext in STYLE_EXT | {".tex"}:
                hits = [(i, l.rstrip()) for i, l in enumerate(txt.splitlines(), 1)
                        if DIGEST_RE.search(l) and not l.lstrip().startswith("%")]
                if hits:
                    digest.append(f"\n## {f.name}  ({len(hits)} lines" +
                                  (f", showing {MAX_DIGEST_LINES_PER_FILE}" if len(hits) > MAX_DIGEST_LINES_PER_FILE else "") + ")\n")
                    digest += [f"{f.name}:{i}: {l[:200]}" for i, l in hits[:MAX_DIGEST_LINES_PER_FILE]]
        elif ext not in (".png", ".jpg", ".jpeg", ".gif", ".eps", ".svg"):
            rec.setdefault("note", "not converted")
        inventory.append(rec)

    (out / "style_digest.md").write_text("\n".join(digest), encoding="utf-8")
    (out / "inventory.json").write_text(json.dumps(inventory, indent=2, ensure_ascii=False), encoding="utf-8")

    print(f"{len(inventory)} files -> {out}/")
    for r in inventory:
        flag = "T" if "text_file" in r else "-"
        print(f"  [{flag}] {r['name']:<40} {r['bytes']:>9} B  {r['role']}")
    print(f"\nRead {out}/style_digest.md and the .txt files next. Hashes for official style files are in inventory.json.")


if __name__ == "__main__":
    sys.exit(main())
